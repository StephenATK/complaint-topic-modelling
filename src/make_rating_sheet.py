"""
make_rating_sheet.py
Builds a BLIND rating sheet for the "topic interpretability" criterion.

All 22 topics (11 from NMF, 11 from LDA) are shuffled and given codes T01 to T22, so
raters can't tell which model a topic came from. Each row shows the topic's top 10
words and its 3 most representative complaints. The key that maps codes back to
models is saved separately and should NOT be shared with raters.

Usage:
    pip install openpyxl          (once, if it isn't installed)
    python src/make_rating_sheet.py

Writes: ratings/rating_sheet_TEMPLATE.xlsx   give a copy to each group member
        data/rating_key.csv                   keep this to yourself until rating is done
"""

import os
import random
import re
import sys

import joblib
import numpy as np
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation

PROJECT_ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
os.chdir(PROJECT_ROOT)

N_WORDS, N_EXAMPLES, MAX_CHARS, SEED = 10, 3, 350, 2026
FONT, BOLD = Font(name="Arial", size=10), Font(name="Arial", size=10, bold=True)
WHITE_BOLD = Font(name="Arial", size=10, bold=True, color="FFFFFF")
NAVY = PatternFill("solid", fgColor="1B3865")
YELLOW = PatternFill("solid", fgColor="FFF2A8")     # cells raters fill in
GREY = PatternFill("solid", fgColor="EEF2F8")
THIN = Side(style="thin", color="C3CFDF")
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
WRAP = Alignment(wrap_text=True, vertical="top")


def clean(text):
    """Readable version of a CFPB narrative for raters."""
    text = re.sub(r"\bX{2,}(?:[/\-]X{2,})*\b", "[redacted]", str(text))
    text = re.sub(r"\s+", " ", text).strip()
    return text if len(text) <= MAX_CHARS else text[:MAX_CHARS].rsplit(" ", 1)[0] + " ..."


def topics_for(model_name, model, vectorizer, doc_topic, texts):
    words = vectorizer.get_feature_names_out()
    rows = []
    for t in range(model.components_.shape[0]):
        top_words = [words[i] for i in model.components_[t].argsort()[::-1][:N_WORDS]]
        best_docs = np.argsort(doc_topic[:, t])[::-1][:N_EXAMPLES]
        rows.append({"model": model_name, "topic": t, "words": ", ".join(top_words),
                     "examples": [clean(texts[d]) for d in best_docs]})
    return rows


def main():
    df = pd.read_csv("data/complaints_clean.csv")
    texts = df["raw_text"].astype(str).tolist() if "raw_text" in df.columns else df["clean_text"].astype(str).tolist()
    topics = []
    for name, model_file, vec_file in [("NMF", "nmf_model", "tfidf_vectorizer"), ("LDA", "lda_model", "count_vectorizer")]:
        doc_topic = np.load(f"data/{name.lower()}_doc_topic.npy")
        if doc_topic.shape[0] != len(texts):
            sys.exit(f"{name}: {doc_topic.shape[0]} topic rows but {len(texts)} complaints. Re-run modeling.py first.")
        topics += topics_for(name, joblib.load(f"data/{model_file}.joblib"), joblib.load(f"data/{vec_file}.joblib"), doc_topic, texts)

    random.Random(SEED).shuffle(topics)
    for i, t in enumerate(topics, 1):
        t["code"] = f"T{i:02d}"
    os.makedirs("ratings", exist_ok=True)
    pd.DataFrame([{"code": t["code"], "model": t["model"], "topic": t["topic"]} for t in topics]).to_csv("data/rating_key.csv", index=False)

    wb = Workbook()
    # ---------- Sheet 1: instructions, legend and a worked example ----------
    ws = wb.active; ws.title = "How to rate"
    lines = [
        ("Topic interpretability rating", Font(name="Arial", size=14, bold=True, color="1B3865")),
        ("Save a copy of this file as ratings/<your name>.xlsx, then rate every row on the Ratings sheet.", FONT),
        ("Work alone. Don't compare scores with other raters until everyone has finished.", FONT),
        ("", FONT),
        ("Fill in only the yellow cells. Each row is one topic. The model it came from is hidden on purpose.", BOLD),
        ("Word clarity (1 to 5): do the top 10 words point to one clear theme?", FONT),
        ("Example fit (1 to 5): do the 3 example complaints match that theme?", FONT),
        ("Suggested name: a short name for the topic, or leave blank if you can't name it.", FONT),
        ("Overall is calculated for you.", FONT),
        ("", FONT),
        ("Scale: 1 = no clear theme   2 = weak   3 = partly clear   4 = clear   5 = very clear", BOLD),
        ("", FONT),
        ("Worked example (not part of the rating):", BOLD),
    ]
    for r, (text, font) in enumerate(lines, 1):
        ws.cell(row=r, column=1, value=text).font = font
        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=7)   # one line across the sheet
    head = ["Code", "Top 10 words", "Example complaint 1", "Word clarity (1-5)", "Example fit (1-5)", "Overall", "Suggested name"]
    example = ["EX", "late, fee, charge, statement, interest, month, added, payment, balance, waived",
               "I was charged a late fee even though my payment was sent before the due date.", 4, 3, None, "Late fees"]
    hr = len(lines) + 1
    for c, h in enumerate(head, 1):
        cell = ws.cell(row=hr, column=c, value=h); cell.font = WHITE_BOLD; cell.fill = NAVY; cell.border = BOX; cell.alignment = WRAP
    for c, v in enumerate(example, 1):
        cell = ws.cell(row=hr + 1, column=c, value=v); cell.font = FONT; cell.border = BOX; cell.alignment = WRAP
    ws.cell(row=hr + 1, column=6, value=f"=AVERAGE(D{hr + 1}:E{hr + 1})").font = FONT
    for c, w in zip("ABCDEFG", [8, 40, 45, 14, 14, 10, 18]):
        ws.column_dimensions[c].width = w

    # ---------- Sheet 2: the 22 topics ----------
    rs = wb.create_sheet("Ratings")
    cols = ["Code", "Top 10 words", "Example complaint 1", "Example complaint 2", "Example complaint 3",
            "Word clarity (1-5)", "Example fit (1-5)", "Overall", "Suggested name", "Notes"]
    widths = [7, 30, 42, 42, 42, 12, 12, 9, 20, 24]
    for c, (h, w) in enumerate(zip(cols, widths), 1):
        cell = rs.cell(row=1, column=c, value=h); cell.font = WHITE_BOLD; cell.fill = NAVY; cell.border = BOX; cell.alignment = WRAP
        rs.column_dimensions[cell.column_letter].width = w
    dv = DataValidation(type="whole", operator="between", formula1="1", formula2="5", allow_blank=True,
                        showErrorMessage=True, errorTitle="Rating", error="Enter a whole number from 1 to 5.")
    rs.add_data_validation(dv)
    for i, t in enumerate(topics, 2):
        values = [t["code"], t["words"]] + t["examples"] + [None, None, None, None, None]
        for c, v in enumerate(values, 1):
            cell = rs.cell(row=i, column=c, value=v); cell.font = FONT; cell.border = BOX; cell.alignment = WRAP
        rs.cell(row=i, column=1).font = BOLD
        rs.cell(row=i, column=8, value=f'=IF(COUNT(F{i}:G{i})=2,AVERAGE(F{i}:G{i}),"")').fill = GREY
        for c in (6, 7, 9, 10):
            rs.cell(row=i, column=c).fill = YELLOW
        dv.add(f"F{i}:G{i}")
        rs.row_dimensions[i].height = 150
    last = len(topics) + 1
    rs.cell(row=last + 2, column=2, value="Rows rated").font = BOLD
    rs.cell(row=last + 2, column=3, value=f"=COUNT(F2:F{last})&\" of {len(topics)}\"").font = FONT
    rs.freeze_panes = "B2"
    for sheet in (ws, rs):                       # print landscape, one page wide
        sheet.page_setup.orientation = "landscape"
        sheet.sheet_properties.pageSetUpPr.fitToPage = True
        sheet.page_setup.fitToWidth, sheet.page_setup.fitToHeight = 1, 0
    rs.print_title_rows = "1:1"
    wb.active = 1
    wb.save("ratings/rating_sheet_TEMPLATE.xlsx")
    print(f"Wrote ratings/rating_sheet_TEMPLATE.xlsx ({len(topics)} topics) and data/rating_key.csv")
    print("Give each group member a copy of the template. Don't share the key.")


if __name__ == "__main__":
    main()
